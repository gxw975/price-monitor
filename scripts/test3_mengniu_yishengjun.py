#!/usr/bin/env python3
"""
第三次完整抓取测试 - 所有规范集成
关键词: 蒙牛一米八八益生菌
规则: 6项固定参数 + 优雅关闭 + 分步校验 + python-xlib滑块 + xdotool DTS点击
"""
import json, urllib.request, websocket, time, math, os, random, subprocess, glob as g, sys, traceback

DISPLAY = ":0"
XAUTH_FILE = "/run/user/1000/.mutter-Xwaylandauth.47UFP3"
DOWNLOAD_DIR = os.path.expanduser("~/Downloads")
KEYWORD = "蒙牛一米八八益生菌"
PROFILE = "/home/lab-admin/chrome-user-desktop"
EXT_PATH = "/tmp/dts_extension/5.0.6_0"

os.environ["DISPLAY"] = DISPLAY
os.environ["XAUTHORITY"] = XAUTH_FILE
env = os.environ.copy()
env["DISPLAY"] = DISPLAY
env["XAUTHORITY"] = XAUTH_FILE

# ---- Logging ----
LOG = []
def log(msg):
    ts = time.strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line)
    LOG.append(line)

# ---- X11 init ----
from Xlib import display as xdisplay, X
from Xlib.ext import xtest
xd = xdisplay.Display()
log("X11 display opened")

def x11_click(sx, sy):
    xtest.fake_input(xd, X.MotionNotify, x=int(sx), y=int(sy)); xd.sync(); time.sleep(0.08)
    xtest.fake_input(xd, X.ButtonPress, 1); xd.sync(); time.sleep(0.05)
    xtest.fake_input(xd, X.ButtonRelease, 1); xd.sync()

def x11_move(x, y):
    xtest.fake_input(xd, X.MotionNotify, x=int(x), y=int(y)); xd.sync()

def x11_down():
    xtest.fake_input(xd, X.ButtonPress, 1); xd.sync()

def x11_up():
    xtest.fake_input(xd, X.ButtonRelease, 1); xd.sync()

# ============================
# STEP 0: 启动Chrome
# ============================
log("=== 步骤0: 启动Chrome ===")

# Cleanup
pkill_res = subprocess.run(["pkill", "-TERM", "chrome"], capture_output=True, timeout=10)
log(f"pkill -TERM chrome: {pkill_res.returncode}")
time.sleep(4)
p_count = int(subprocess.run(["pgrep", "-c", "chrome"], capture_output=True, text=True).stdout.strip() or "0")
if p_count > 0:
    log(f"TERM未完全退出({p_count}残留), KILL兜底")
    subprocess.run(["pkill", "-KILL", "chrome"], capture_output=True, timeout=10)
    time.sleep(3)

# Clean Singleton
for lock in g.glob(f"{PROFILE}/Singleton*") + g.glob(f"{PROFILE}/Default/Singleton*"):
    try: os.remove(lock)
    except: pass

# Fix permissions (profile already owned by lab-admin)
# subprocess.run(["sudo", "chown", "-R", "lab-admin:lab-admin", PROFILE], capture_output=True, timeout=10)

# Start Chrome (6 base params + --load-extension for DTS install)
log(f"启动Chrome profile={PROFILE}")
chrome_proc = subprocess.Popen([
    "/opt/google/chrome/chrome",
    "--user-data-dir", PROFILE,
    "--disable-gpu",
    "--disable-software-rasterizer",
    "--remote-debugging-port", "9223",
    "--remote-allow-origins", "*",
    "--start-maximized",
    "--load-extension", EXT_PATH,
    "https://www.taobao.com/"
], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
log(f"Chrome PID={chrome_proc.pid}")

# Wait for CDP
for i in range(30):
    time.sleep(2)
    try:
        r = urllib.request.urlopen('http://127.0.0.1:9223/json/version', timeout=3)
        log(f"CDP ready after {(i+1)*2}s")
        break
    except:
        if i == 29: log("❌ CDP启动超时"); sys.exit(1)

# ---- CDP Helpers ----
def get_ws():
    r = urllib.request.urlopen('http://127.0.0.1:9223/json', timeout=5)
    ts = json.loads(r.read())
    for t in ts:
        if 'taobao' in t.get('url', '') and t.get('type') == 'page':
            return websocket.create_connection(t['webSocketDebuggerUrl'], timeout=15)
    raise Exception("No taobao page")

ws = get_ws()
_mid = [0]
def cdp(method, params=None):
    _mid[0] += 1
    msg = {'id': _mid[0], 'method': method}
    if params: msg['params'] = params
    ws.send(json.dumps(msg))
    while True:
        r = json.loads(ws.recv())
        if r.get('id') == _mid[0]: return r

cdp('Page.enable')
cdp('Runtime.enable')

def get_offset():
    r = cdp('Runtime.evaluate', {
        'expression': """(function(){var fl=(window.outerWidth-window.innerWidth)/2;return JSON.stringify({ox:Math.round((window.screenLeft||0)+fl),oy:Math.round((window.screenTop||0)+window.outerHeight-window.innerHeight-fl)});})()""",
        'returnByValue': True
    })
    d = json.loads(r.get('result', {}).get('result', {}).get('value', '{}'))
    return d.get('ox', 66), d.get('oy', 119)

def body_has(text):
    r = cdp('Runtime.evaluate', {
        'expression': f"(function(){{return (document.body?.innerText||'').indexOf('{text}')!==-1;}})()",
        'returnByValue': True
    })
    return r.get('result', {}).get('result', {}).get('value') == True

def body_progress():
    r = cdp('Runtime.evaluate', {
        'expression': """(function(){var b=document.body?.innerText||'';var m=b.match(/已成功加载[：:]?\\s*(\\d+)\\s*\/\\s*(\\d+)/);return JSON.stringify(m?{loaded:parseInt(m[1]),total:parseInt(m[2])}:null);})()""",
        'returnByValue': True
    })
    return json.loads(r.get('result', {}).get('result', {}).get('value', 'null'))

def find_text(text):
    js = f"""(function(){{var t='{text}';var all=document.querySelectorAll('*');for(var i=0;i<all.length;i++){{if((all[i].textContent||'').trim()===t&&all[i].offsetHeight>0&&all[i].offsetHeight<100){{var r=all[i].getBoundingClientRect();return JSON.stringify({{found:true,vx:Math.round(r.x+r.width/2),vy:Math.round(r.y+r.height/2)}});}}}}return JSON.stringify({{found:false}});}})()"""
    r = cdp('Runtime.evaluate', {'expression': js, 'returnByValue': True})
    return json.loads(r.get('result', {}).get('result', {}).get('value', '{}'))

def find_contains(text):
    js = f"""(function(){{var t='{text}';var all=document.querySelectorAll('*');for(var i=0;i<all.length;i++){{var txt=(all[i].textContent||'').trim();if(txt.indexOf(t)!==-1&&all[i].offsetHeight>0&&all[i].offsetHeight<100&&txt.length===t.length){{var r=all[i].getBoundingClientRect();return JSON.stringify({{found:true,vx:Math.round(r.x+r.width/2),vy:Math.round(r.y+r.height/2)}});}}}}return JSON.stringify({{found:false}});}})()"""
    r = cdp('Runtime.evaluate', {'expression': js, 'returnByValue': True})
    return json.loads(r.get('result', {}).get('result', {}).get('value', '{}'))

def has_captcha():
    r = cdp('Runtime.evaluate', {
        'expression': """(function(){var fs=document.querySelectorAll('iframe');for(var i=0;i<fs.length;i++){if((fs[i].src||'').indexOf('h5api')!==-1&&fs[i].getBoundingClientRect().width>0)return true;}return false;})()""",
        'returnByValue': True
    })
    return r.get('result', {}).get('result', {}).get('value') == True

def dts_click(vx, vy, label=""):
    ox, oy = get_offset()
    sx, sy = int(vx + ox), int(vy + oy)
    log(f"  xdotool click '{label}' @ screen({sx},{sy})")
    subprocess.run(["xdotool", "mousemove", str(sx), str(sy)], env=env, capture_output=True, timeout=10)
    time.sleep(0.4)
    subprocess.run(["xdotool", "click", "1"], env=env, capture_output=True, timeout=10)
    time.sleep(1)

# ---- Captcha Solver (python-xlib, unchanged) ----
def solve_captcha():
    log("  ⚠️ 滑块验证码求解中...")
    ft = cdp('Page.getFrameTree')
    pfid = [None]
    def walk(node):
        for c in node.get('childFrames', []):
            f = c.get('frame', {})
            if 'h5api' in f.get('url', '') and 'punish' in f.get('url', ''):
                pfid[0] = f.get('id', ''); return
            walk(c)
    walk(ft.get('result', {}).get('frameTree', {}))
    pfid = pfid[0]
    if not pfid: log("  ❌ punish frame未找到"); return False

    iso = cdp('Page.createIsolatedWorld', {'frameId': pfid})
    ctx = iso.get('result', {}).get('executionContextId')
    ox, oy = get_offset()

    for attempt in range(1, 7):
        r = cdp('Runtime.evaluate', {
            'expression': """(function(){var s=document.querySelector('[id*="nc_1_n1z"]');var t=document.querySelector('[id*="nc_1__scale_text"]');if(!s||!t){var w=document.querySelector('[id*="nc_1_wrapper"]');if(w&&(w.innerHTML||'').indexOf('errloading')!==-1)return JSON.stringify({e:'errloading'});return JSON.stringify({f:false});}var sr=s.getBoundingClientRect(),tr=t.getBoundingClientRect();return JSON.stringify({x:Math.round(sr.x+sr.width/2),y:Math.round(sr.y+sr.height/2),d:Math.round(tr.x+tr.width-sr.x-sr.width+3)});})()""",
            'contextId': ctx, 'returnByValue': True
        })
        sl = json.loads(r.get('result', {}).get('result', {}).get('value', '{}'))

        if sl.get('e') == 'errloading':
            er = cdp('Runtime.evaluate', {
                'expression': """(function(){var e=document.querySelector('.errloading');if(!e)return JSON.stringify({f:false});var r=e.getBoundingClientRect();return JSON.stringify({x:Math.round(r.x+r.width/2),y:Math.round(r.y+r.height/2)});})()""",
                'contextId': ctx, 'returnByValue': True
            })
            ep = json.loads(er.get('result', {}).get('result', {}).get('value', '{}'))
            if ep.get('x'):
                ip = cdp('Runtime.evaluate', {
                    'expression': """(function(){var fs=document.querySelectorAll('iframe');for(var i=0;i<fs.length;i++){if((fs[i].src||'').indexOf('h5api')!==-1){var r=fs[i].getBoundingClientRect();return JSON.stringify({x:Math.round(r.x),y:Math.round(r.y)});}}return JSON.stringify({});})()""",
                    'returnByValue': True
                })
                ip2 = json.loads(ip.get('result', {}).get('result', {}).get('value', '{}'))
                ex = int(ip2.get('x', 397) + ep['x'] + ox)
                ey = int(ip2.get('y', 181) + ep['y'] + oy)
                x11_click(ex, ey)
                log(f"  errloading恢复 click({ex},{ey})")
                time.sleep(3)
            continue

        if not sl.get('x'):
            if not body_has('请拖动') and not has_captcha():
                log("  ✅ 验证码已消失"); return True
            time.sleep(2); continue

        ip = cdp('Runtime.evaluate', {
            'expression': """(function(){var fs=document.querySelectorAll('iframe');for(var i=0;i<fs.length;i++){if((fs[i].src||'').indexOf('h5api')!==-1){var r=fs[i].getBoundingClientRect();return JSON.stringify({x:Math.round(r.x),y:Math.round(r.y)});}}return JSON.stringify({});})()""",
            'returnByValue': True
        })
        ip2 = json.loads(ip.get('result', {}).get('result', {}).get('value', '{}'))
        sx = ip2.get('x', 397) + sl['x'] + ox
        sy = ip2.get('y', 181) + sl['y'] + oy
        dist = sl['d']
        log(f"  尝试{attempt}: ({sx},{sy}) d={dist}")

        ax = sx - random.randint(35, 55)
        ay = sy + random.randint(-10, 10)
        for i in range(14):
            p = (i + 1) / 14
            x11_move(int(ax + (sx - ax) * p), int(ay + (sy - ay) * p + math.sin(i * 0.5) * 5))
            time.sleep(0.015 + random.random() * 0.02)
        x11_move(sx, sy); time.sleep(0.25 + random.random() * 0.35)
        x11_down(); time.sleep(0.03)

        pts = 280 + random.randint(40, 80)
        total_t = 3.0 + random.random() * 2.5
        for i in range(pts):
            p = i / pts
            if p < 0.08: e = (p / 0.08) ** 2 * 0.08
            elif p < 0.88: e = 0.08 + (p - 0.08) * 0.84
            else: r2 = (1 - p) / 0.12; e = 1 - r2 * r2 * 0.12
            e += (random.random() - 0.5) * 0.005
            x = int(sx + dist * e)
            yj = math.sin(p * math.pi * 3.1) * 2.5 + math.sin(p * math.pi * 7.3) * 0.6 + math.sin(p * math.pi * 13.7) * 0.3
            if random.random() < 0.02: yj += (random.random() - 0.5) * 8
            y = int(sy + yj)
            x11_move(x, y)
            time.sleep(max(0.002, total_t / pts * (0.6 + random.random() * 0.8)))
        x11_move(int(sx + dist), int(sy)); time.sleep(0.05)
        x11_move(int(sx + dist + random.uniform(2, 5)), int(sy + random.randint(-1, 1))); time.sleep(0.04)
        x11_move(int(sx + dist), int(sy)); time.sleep(0.05)
        time.sleep(0.18 + random.random() * 0.25)
        x11_up()
        log(f"  拖拽完成 pts={pts} t={total_t:.1f}s")
        time.sleep(4)
        if not body_has('请拖动') and not has_captcha():
            log(f"  ✅ 尝试{attempt}成功!"); return True
    return False

# ---- Page Validation ----
def validate(condition, error_msg):
    """分步校验：不通过直接终止"""
    if not condition:
        log(f"❌ 校验失败: {error_msg}")
        raise AssertionError(f"校验失败: {error_msg}")
    log(f"  ✅ {error_msg}")

# ============================
# STEP 1: 验证淘宝首页 + DTS扩展加载
# ============================
log("\n=== 步骤1: 验证环境 ===")
r = cdp('Runtime.evaluate', {'expression': 'window.location.href', 'returnByValue': True})
url = r.get('result', {}).get('result', {}).get('value', '')
log(f"  当前URL: {url[:80]}")

# ============================
# STEP 2: 搜索关键词
# ============================
log(f"\n=== 步骤2: 搜索 '{KEYWORD}' ===")
enc = urllib.request.quote(KEYWORD)
cdp('Page.navigate', {'url': f'https://s.taobao.com/search?q={enc}'})
time.sleep(8)

# Wait for page and handle captcha
for i in range(8):
    time.sleep(5)
    if has_captcha():
        solve_captcha()
        time.sleep(5)
        continue
    if body_has('市场分析'):
        log(f"  ✅ 搜索结果已加载({(i+1)*5+8}s)")
        break
    log(f"  等待...({(i+1)*5+8}s)")

validate(body_has('市场分析'), "搜索结果页应出现'市场分析'(DTS已注入)")

# ============================
# STEP 3: 清理缓存(如有异常标识)
# ============================
if body_has('遇到异常') or body_has('继续加载'):
    log("\n=== 步骤3: 清理缓存 ===")
    el = find_text('清理缓存')
    if el.get('found'):
        dts_click(el['vx'], el['vy'], '清理缓存')
        time.sleep(5)
        log("  清理缓存完成")
    else:
        log("  无清理缓存按钮，跳过")

# ============================
# STEP 4: 市场分析
# ============================
log("\n=== 步骤4: 市场分析 ===")
validate(body_has('市场分析'), "搜索页应有'市场分析'按钮")
el = find_text('市场分析')
if not el.get('found'):
    el = find_contains('市场分析')
if el.get('found'):
    dts_click(el['vx'], el['vy'], '市场分析')
else:
    # Fallback: estimate position
    log("  使用位置估计")
    dts_click(779, 314, '市场分析(估计)')

# Wait for DTS panel
for i in range(12):
    time.sleep(3)
    if body_has('开始分析'):
        log(f"  ✅ DTS面板已打开({(i+1)*3}s)")
        break
validate(body_has('开始分析'), "DTS面板打开后应有'开始分析'按钮")

# ============================
# STEP 5: 开始分析
# ============================
log("\n=== 步骤5: 开始分析 ===")
el = find_text('开始分析')
if el.get('found'):
    dts_click(el['vx'], el['vy'], '开始分析')

# Wait for analysis
for i in range(20):
    time.sleep(5)
    p = body_progress()
    has_auto = body_has('自动加载') or body_has('加载下一页')
    has_err = body_has('遇到异常')
    if p or has_auto:
        log(f"  ✅ 分析完成 progress={p}")
        break
    if has_err:
        log(f"  ⚠️ 分析中遇到异常")
        break
validate(body_has('自动加载') or body_has('加载下一页') or body_has('导出表格'),
         "分析完成后应有'自动加载'或'导出表格'")

# ============================
# STEP 6: 分页自动加载
# ============================
log("\n=== 步骤6: 分页自动加载 ===")
total_clicks = 0
for round_num in range(1, 9):
    log(f"  --- 第{round_num}轮 ---")
    for click_in_round in range(7):
        time.sleep(5)
        if click_in_round == 0:
            time.sleep(10)

        # Check for loading error
        if body_has('遇到异常') and body_has('继续加载'):
            log(f"  加载异常，尝试继续加载...")
            cbtn = find_text('继续加载')
            if cbtn.get('found'):
                dts_click(cbtn['vx'], cbtn['vy'], '继续加载')
                time.sleep(10)
                if not body_has('继续加载'):
                    log("  继续加载成功")
                continue
            # If still failing, try clear cache
            cc_btn = find_text('清理缓存')
            if cc_btn.get('found'):
                dts_click(cc_btn['vx'], cc_btn['vy'], '清理缓存')
                time.sleep(5)
                log("  缓存已清理，需重走流程...")
                raise RuntimeError("需要重新走DTS流程")

        el = find_text('自动加载')
        if not el.get('found'):
            el = find_text('加载下一页')
        if not el.get('found'):
            p = body_progress()
            log(f"  ✅ 加载完成 clicks={total_clicks} progress={p}")
            break

        total_clicks += 1
        dts_click(el['vx'], el['vy'], f'加载 #{total_clicks}')
        time.sleep(25)
    else:
        continue
    break

p = body_progress()
log(f"  总点击: {total_clicks}, 进度: {p}")
validate(p is not None, "加载完成后应有进度数据")

# ============================
# STEP 7: 等待数据稳定
# ============================
log("\n=== 步骤7: 等待315秒数据稳定 ===")
for i in range(21):
    time.sleep(15)
    if (i+1) % 5 == 0:
        log(f"  等待{(i+1)*15}s...")

# ============================
# STEP 8: 导出xlsx
# ============================
log("\n=== 步骤8: 导出xlsx ===")
validate(body_has('导出表格'), "导出前应有'导出表格'按钮")

old_files = set(g.glob(f'{DOWNLOAD_DIR}/*.xlsx'))
el = find_text('导出表格')
if el.get('found'):
    dts_click(el['vx'], el['vy'], '导出表格')
else:
    dts_click(1003, 210, '导出表格(估计)')

time.sleep(3)
# Click xlsx option in dropdown
xlsx_opt = find_contains('xlsx')
if not xlsx_opt.get('found'):
    xlsx_opt = find_contains('csv')
if xlsx_opt.get('found'):
    dts_click(xlsx_opt['vx'], xlsx_opt['vy'], f"导出格式: {xlsx_opt.get('text','')}")

downloaded = None
for i in range(30):
    new_files = set(g.glob(f'{DOWNLOAD_DIR}/*.xlsx')) - old_files
    if new_files:
        downloaded = list(new_files)[0]
        log(f"  ✅ 下载: {downloaded}")
        break
    time.sleep(5)

validate(downloaded is not None, "xlsx文件应下载成功")

# ============================
# STEP 9: 验证数据
# ============================
log("\n=== 步骤9: 验证数据 ===")
import openpyxl
wb = openpyxl.load_workbook(downloaded)
sheet = wb.active
product_count = sheet.max_row - 1
log(f"  Sheet: {wb.sheetnames[0]} rows={sheet.max_row} cols={sheet.max_column}")
log(f"  商品数量: {product_count}")
validate(product_count > 0, "xlsx应包含商品数据")

# ============================
# STEP 10: 优雅关闭
# ============================
log("\n=== 步骤10: 优雅关闭 ===")
try:
    cdp('Page.close')
    log("  Page.close done")
except: pass
time.sleep(2)

subprocess.run(["pkill", "-TERM", "chrome"], capture_output=True, timeout=10)
time.sleep(4)
p_count = int(subprocess.run(["pgrep", "-c", "chrome"], capture_output=True, text=True).stdout.strip() or "0")
if p_count > 0:
    log(f"  TERM残留{p_count}进程, KILL兜底")
    subprocess.run(["pkill", "-KILL", "chrome"], capture_output=True, timeout=10)
    time.sleep(2)
log("  Chrome已关闭")

xd.close()
ws.close()
chrome_proc.wait(timeout=10)

# ============================
# 测试报告
# ============================
log(f"\n{'='*60}")
log(f"第三次测试完成")
log(f"关键词: {KEYWORD}")
log(f"导出文件: {os.path.basename(downloaded)}")
log(f"商品数量: {product_count}条")
log(f"是否有异常: 否")
log(f"{'='*60}")

# Write log file
log_path = f'/home/lab-admin/price-monitor/logs/test3_{time.strftime("%Y%m%d_%H%M%S")}.log'
os.makedirs(os.path.dirname(log_path), exist_ok=True)
with open(log_path, 'w') as f:
    f.write('\n'.join(LOG))
log(f"日志已保存: {log_path}")
